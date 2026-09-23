#!/usr/bin/env python3
"""The routine edition of the tour plan, as an Excel workbook.

    python3 tools/build-plan-routine-xlsx.py plan/tour-plan-<stay>.json plan/tour-plan-<stay>-routine.xlsx

One tab per day with only the travel and event routine: from leaving the
stay to leaving the last venue. No wake-up / breakfast / getting ready at
the stay, no return leg or lights-out. A dinner is always there before the
group leaves the last venue: when the plan has dinner back at the stay, a
dinner near the venue is added at the end of the day's routine (marked so).
Arrival and departure days are included (land → stay; dinner → airport →
flight). Instrument-vehicle rows are left out to keep it simple.
"""
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK, MUTED, COPPER, CREAM, LINE, WHITE, AMBER = "1B1D21", "666666", "8E4A24", "F7F3EA", "D9D4C8", "FFFFFF", "6B4A00"
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DINNER_MIN = 60


def hm(m):
    if m is None:
        return ""
    m = int(round(m)); suffix = ""
    if m >= 1440:
        m -= 1440; suffix = " (+1)"
    if m < 0:
        m += 1440; suffix = " (−1)"
    h, mm = divmod(m, 60)
    return f"{12 if h % 12 == 0 else h % 12}.{mm:02d}{'am' if h < 12 else 'pm'}{suffix}"


def day_label(iso):
    d = date.fromisoformat(iso)
    return f"{DAYS[d.weekday()]}, {d.day} {MONTHS[d.month - 1]}"


def tab_name(iso):
    d = date.fromisoformat(iso)
    return f"{d.day:02d} {MONTHS[d.month - 1]}"


thin = Side(style="thin", color=LINE)
BORDER = Border(bottom=thin)
F_TITLE = Font(name="Arial", size=16, bold=True, color=INK)
F_KICK = Font(name="Arial", size=9, bold=True, color=COPPER)
F_SUB = Font(name="Arial", size=10, color=MUTED)
F_HEAD = Font(name="Arial", size=9, bold=True, color=WHITE)
F_BODY = Font(name="Arial", size=10, color=INK)
F_BOLD = Font(name="Arial", size=10, bold=True, color=INK)
F_MUTED = Font(name="Arial", size=9, color=MUTED)
F_TIME = Font(name="Arial", size=10, color=COPPER)
F_SHOW = Font(name="Arial", size=11, bold=True, color=WHITE)
F_SHOW_SUB = Font(name="Arial", size=9, color="D5D0C5")
F_MEAL = Font(name="Arial", size=10, bold=True, color=COPPER)
F_NOTE = Font(name="Arial", size=9, italic=True, color=AMBER)
FILL_HEAD = PatternFill("solid", fgColor=INK)
FILL_SHOW = PatternFill("solid", fgColor=INK)
FILL_LEG = PatternFill("solid", fgColor=CREAM)
WRAP = Alignment(wrap_text=True, vertical="top")


def is_stay_return(b):
    l = (b.get("label") or "").lower()
    return b.get("type") == "leg" and (l.startswith("return to the stay") or l.startswith("back to the stay"))


def routine(day):
    """The blocks that make up the day's routine, plus a note when a dinner
    was added."""
    blocks = [b for b in day.get("blocks") or [] if b.get("type") != "truck"]
    kind = day.get("kind")
    note = ""
    if kind == "rest":
        return [], "Rest day at the stay · no travel"
    if kind == "arrival":
        keep = [b for b in blocks if b.get("type") in ("show", "leg", "free", "meal")]
        return keep, ""
    if kind == "departure":
        start = next((i for i, b in enumerate(blocks) if b.get("meal") == "dinner"), None)
        if start is None:
            start = next((i for i, b in enumerate(blocks) if b.get("type") == "leg"), 0)
        return blocks[start:], ""
    # Show / outstation days: from the first departure to the last departure
    # from a venue.
    first_leg = next((i for i, b in enumerate(blocks) if b.get("type") == "leg"), 0)
    keep = blocks[first_leg:]
    # Drop the final return to the stay and anything after it.
    last_return = max((i for i, b in enumerate(keep) if is_stay_return(b)), default=None)
    if last_return is not None and last_return == max((i for i, b in enumerate(keep) if b.get("type") == "leg"), default=-1):
        keep = keep[:last_return]
    # Drop trailing broad blocks (sleep) and anything at the stay after the last leg.
    while keep and (keep[-1].get("type") in ("sleep", "wake") or (keep[-1].get("broad") and keep[-1].get("meal") != "dinner")):
        keep.pop()
    # A dinner before the group leaves the last venue.
    if keep and not any(b.get("meal") == "dinner" for b in keep):
        last = keep[-1]
        at = last.get("to") if last.get("to") is not None else last.get("from")
        venue = ""
        for b in reversed(keep):
            if b.get("type") == "show" and b.get("venue"):
                venue = b["venue"].split(",")[0]; break
        keep.append({"type": "meal", "meal": "dinner", "from": at, "to": (at or 0) + DINNER_MIN,
                     "label": "Dinner near " + venue if venue else "Dinner before leaving", "detail": "added: the plan has dinner back at the stay", "added": True})
        note = "Dinner added before leaving the venue; the plan itself has dinner at the stay after the drive back."
    return keep, note


def build(plan, out):
    wb = Workbook()
    stay = plan.get("stay") or {}
    party = plan.get("party") or {}
    days = plan.get("days") or []
    ws = wb.active; ws.title = "Days"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "BALI IN BENGALURU · ARTIST TOUR · DAILY ROUTINE"; ws["A1"].font = F_KICK
    ws["A2"] = "Travel and event routine, day by day"; ws["A2"].font = F_TITLE
    ws["A3"] = f"{party.get('artists', '—')} artists + {party.get('volunteers', 0)} volunteers · staying at {stay.get('name', '—')} · from leaving the stay to leaving the last venue, dinner included · one tab per day"; ws["A3"].font = F_SUB
    for i, (label, w) in enumerate([("Date", 14), ("Day", 12), ("Events", 46), ("Leave the stay", 14), ("Last venue", 34), ("Dinner", 22), ("Note", 50)], 1):
        c = ws.cell(row=5, column=i, value=label); c.font = F_HEAD; c.fill = FILL_HEAD
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"
    r = 6
    for d in days:
        keep, note = routine(d)
        kind = {"rest": "Rest day", "outstation": "Out of town", "arrival": "Arrival", "departure": "Departure"}.get(d.get("kind"), "Show day")
        evs = "; ".join(e["title"] + (f" {hm(e['start'])}" if e.get("start") is not None else "") for e in d.get("events", []))
        leave = next((b for b in keep if b.get("type") == "leg"), None)
        last_show = next((b for b in reversed(keep) if b.get("type") == "show"), None)
        dinner = next((b for b in keep if b.get("meal") == "dinner"), None)
        row = [day_label(d["date"]), kind, evs, hm(leave["from"]) if leave else "", (last_show or {}).get("venue", "").split(",")[0] if last_show else "",
               (hm(dinner["from"]) + " · " + dinner["label"]) if dinner else "", note or " · ".join(d.get("flags") or [])[:200]]
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v); c.font = F_BODY; c.alignment = WRAP; c.border = BORDER
        ws.cell(row=r, column=1).font = F_BOLD
        if d.get("kind") == "rest":
            for i in range(1, 8): ws.cell(row=r, column=i).font = F_MUTED
        r += 1

    for d in days:
        keep, note = routine(d)
        ws = wb.create_sheet(tab_name(d["date"]))
        ws.sheet_view.showGridLines = False
        ws["A1"] = "BALI IN BENGALURU · DAILY ROUTINE"; ws["A1"].font = F_KICK
        ws["A2"] = day_label(d["date"]); ws["A2"].font = F_TITLE
        evs = " · ".join(e["title"] + (f" {hm(e['start'])}" if e.get("start") is not None else "") for e in d.get("events", []))
        ws["A3"] = evs or ({"rest": "Rest day at the stay", "arrival": "Arrival day", "departure": "Departure day"}.get(d.get("kind"), "")); ws["A3"].font = F_SUB
        r = 4
        jn = next((f for f in d.get("flags") or [] if "perform with the group" in f or "travel with the group" in f), None)
        if jn:
            ws.cell(row=r, column=1, value="• " + jn).font = F_NOTE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); r += 1
        if note:
            ws.cell(row=r, column=1, value="• " + note).font = F_NOTE; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); r += 1
        r += 1
        for i, (label, w) in enumerate([("From", 12), ("To", 12), ("What", 52), ("Where / detail", 64)], 1):
            c = ws.cell(row=r, column=i, value=label); c.font = F_HEAD; c.fill = FILL_HEAD
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{r + 1}"
        r += 1
        if not keep:
            ws.cell(row=r, column=3, value="Rest day at the stay · no travel").font = F_MUTED
            continue
        for b in keep:
            typ = b.get("type")
            what = b.get("label") or ""
            detail = b.get("detail") or ""
            if typ == "show":
                bits = [x for x in [b.get("category"), b.get("venue")] if x]
                if b.get("artists") is not None or b.get("cast"):
                    bits.append(" · ".join(x for x in [f"{b['artists']} artists" if b.get("artists") is not None else "", b.get("cast") or ""] if x))
                detail = " · ".join(bits)
            if b.get("note"):
                detail = (detail + " · " if detail else "") + b["note"]
            if typ == "leg":
                what = "→ " + what
            cells = [hm(b.get("from")), hm(b.get("to")) if b.get("to") is not None else "", what, detail]
            for i, v in enumerate(cells, 1):
                c = ws.cell(row=r, column=i, value=v); c.font = F_BODY; c.alignment = WRAP; c.border = BORDER
            ws.cell(row=r, column=1).font = F_TIME; ws.cell(row=r, column=2).font = F_TIME
            ws.cell(row=r, column=3).font = F_BOLD; ws.cell(row=r, column=4).font = F_MUTED
            if typ == "show":
                for i in range(1, 5): ws.cell(row=r, column=i).fill = FILL_SHOW
                ws.cell(row=r, column=1).font = Font(name="Arial", size=10, color="F3E9D2"); ws.cell(row=r, column=2).font = Font(name="Arial", size=10, color="F3E9D2")
                ws.cell(row=r, column=3).font = F_SHOW; ws.cell(row=r, column=3).value = what.upper()
                ws.cell(row=r, column=4).font = F_SHOW_SUB; ws.row_dimensions[r].height = 22
            elif typ == "leg":
                for i in range(1, 5): ws.cell(row=r, column=i).fill = FILL_LEG
            elif b.get("meal"):
                ws.cell(row=r, column=3).font = F_MEAL
                if b.get("added"):
                    ws.cell(row=r, column=4).font = F_NOTE
            r += 1
    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor = COPPER if sheet.title == "Days" else "8C8E92"
    wb.save(out)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
    print("wrote", sys.argv[2])
